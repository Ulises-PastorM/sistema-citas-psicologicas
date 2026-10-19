from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from models.sesion_model import Sesion
from repositories.sesiones_repository import (
    crear_sesion,
    obtener_sesiones,
    obtener_sesion_por_id,
    obtener_sesiones_por_cita,
    obtener_historial_observaciones,
    actualizar_sesion,
    eliminar_sesion,
)
from utils.utilities import fecha_a_texto


def service_crear_sesion(sesion: Sesion) -> dict:
    if not sesion.cita_id:
        return {"success": False, "error": "El ID de la cita es obligatorio."}

    try:
        nuevo_id = crear_sesion(sesion)
        return {"success": True, "id_sesion": nuevo_id}
    except Exception as e:
        return {"success": False, "error": f"Error al crear la sesión: {str(e)}"}


def service_obtener_sesiones() -> list[Sesion]:
    try:
        return obtener_sesiones()
    except Exception as e:
        print(f"[sesiones_service] Error al obtener sesiones: {e}")
        return []


def service_obtener_sesion_por_id(id_sesion: int) -> Sesion | None:
    if not id_sesion:
        return None
    try:
        return obtener_sesion_por_id(id_sesion)
    except Exception as e:
        print(f"[sesiones_service] Error al obtener sesión {id_sesion}: {e}")
        return None


def service_obtener_sesiones_por_cita(cita_id: int) -> list[Sesion]:
    if not cita_id:
        return []
    try:
        return obtener_sesiones_por_cita(cita_id)
    except Exception as e:
        print(f"[sesiones_service] Error al obtener sesiones de cita {cita_id}: {e}")
        return []


def service_obtener_historial_observaciones():
    try:
        usuarias_observaciones = []
        historial_observaciones = obtener_historial_observaciones()

        if not historial_observaciones:
            return []
        
        #Para concatenar todas las observaciones de una usuaria
        id_usuaria_actual = historial_observaciones[0].id_usuaria
        nombre_usuaria_actual = historial_observaciones[0].nombre
        observaciones_actual = ""

        for obs in historial_observaciones:
            if obs.observaciones and obs.observaciones.strip():
                if obs.id_usuaria == id_usuaria_actual:
                    observaciones_actual += fecha_a_texto(obs.fecha_sesion) + " " + obs.hora_cita + ": " + obs.observaciones + "\n"
                else:
                    if observaciones_actual:
                        usuarias_observaciones.append([nombre_usuaria_actual, observaciones_actual])
                    id_usuaria_actual = obs.id_usuaria
                    nombre_usuaria_actual = obs.nombre
                    observaciones_actual = fecha_a_texto(obs.fecha_sesion) + " " + obs.hora_cita + ": " + obs.observaciones + "\n"
        
        if observaciones_actual:
            usuarias_observaciones.append([nombre_usuaria_actual, observaciones_actual])

        return usuarias_observaciones

    except Exception as e:
        print(f"[sesiones_service] Error al obtener el historial de observaciones: {e}")
        return []


def service_actualizar_sesion(sesion: Sesion) -> dict:
    if not sesion.id_sesion:
        return {"success": False, "error": "El ID de la sesión es requerido."}

    existente = obtener_sesion_por_id(sesion.id_sesion)
    if not existente:
        return {"success": False, "error": f"No se encontró una sesión con ID {sesion.id_sesion}."}

    try:
        filas = actualizar_sesion(sesion)
        if filas == 0:
            return {"success": False, "error": "No se realizaron cambios."}
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": f"Error al actualizar la sesión: {str(e)}"}


def service_eliminar_sesion(id_sesion: int) -> dict:
    if not id_sesion:
        return {"success": False, "error": "El ID de la sesión es requerido."}

    existente = obtener_sesion_por_id(id_sesion)
    if not existente:
        return {"success": False, "error": f"No se encontró una sesión con ID {id_sesion}."}

    try:
        filas = eliminar_sesion(id_sesion)
        if filas == 0:
            return {"success": False, "error": "No se pudo eliminar la sesión."}
        return {"success": True}
    except Exception as e:
        return {"success": False, "error": f"Error al eliminar la sesión: {str(e)}"}


def service_exportar_observaciones_excel(ruta: str) -> dict:
    try:
        historial = service_obtener_historial_observaciones()

        if not historial:
            return {
                "success": False,
                "error": "No hay observaciones para exportar."
            }

        ruta_archivo = Path(ruta).with_suffix(".xlsx")

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Observaciones"

        hoja.append(["Usuaria", "Observaciones"])

        for nombre, observaciones in historial:
            hoja.append([nombre, observaciones.rstrip("\n")])

            # Guardar ambos valores como texto, incluso si empiezan con "=".
            for celda in hoja[hoja.max_row]:
                celda.data_type = "s"
                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="006B4D"
            )

        hoja.column_dimensions["A"].width = 30
        hoja.column_dimensions["B"].width = 100

        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

        libro.save(ruta_archivo)
        libro.close()

        return {
            "success": True,
            "ruta": str(ruta_archivo),
            "total": len(historial)
        }

    except PermissionError:
        return {
            "success": False,
            "error": (
                "No se pudo guardar el archivo. "
                "Comprueba que no esté abierto en Excel "
                "y que tengas permiso para guardar en esa carpeta."
            )
        }
    except Exception as e:
        return {
            "success": False,
            "error": f"Error al exportar las observaciones: {e}"
        }